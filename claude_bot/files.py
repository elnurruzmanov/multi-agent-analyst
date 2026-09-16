"""Claude foydalanuvchiga fayl yasab bera olishi uchun asbob.

Claude fayl yarata olmaydi — u faqat matn qaytaradi. Shuning uchun unga
`create_file` degan qurol beramiz: model qaysi fayl kerakligini va ichida
nima bo'lishini aytadi, faylni esa shu yerda, bot ichida yig'amiz.

Shu tarzda foydalanuvchi tabiiy yozadi («buni excelga sol»), model o'zi
kerakli qurolni chaqiradi — alohida buyruq o'rganish shart emas.

Modul aiogram'ga bog'liq emas: faqat baytlar qaytaradi, yuborish `main.py`
ning ishi.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass

from claude_bot import pdf

KINDS = ("xlsx", "docx", "pdf", "csv", "txt")

EXTENSIONS = {
    "xlsx": ".xlsx", "docx": ".docx", "pdf": ".pdf", "csv": ".csv", "txt": ".txt",
}

# Bitta javobda nechta fayl yasalishi mumkin — cheksiz bo'lmasin.
MAX_FILES = 5
MAX_ROWS = 5000

_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
_BOLD = re.compile(r"\*\*(?=\S)(.+?)(?<=\S)\*\*", re.S)


@dataclass
class Artifact:
    """Yuborishga tayyor fayl."""

    filename: str
    data: bytes
    kind: str


TOOL = {
    "name": "create_file",
    "description": (
        "Foydalanuvchiga fayl yasab yuborish: Excel, Word, PDF, CSV yoki "
        "oddiy matn. Foydalanuvchi fayl so'raganda (masalan «excel qilib ber», "
        "«word hujjat qil», «jadval qilib yubor») shu qurolni chaqir. "
        "Jadval ko'rinishidagi ma'lumot uchun `rows` dan foydalan — birinchi "
        "qator sarlavha bo'lsin. Matnli hujjat uchun `text` ga markdown yoz. "
        "Fayl foydalanuvchiga avtomatik yuboriladi, sen faqat chaqirasan."
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "kind": {
                "type": "string",
                "enum": list(KINDS),
                "description": "Fayl turi. Jadval uchun xlsx yoki csv.",
            },
            "filename": {
                "type": "string",
                "description": "Kengaytmasiz nom, masalan «qarindoshlar_royxati».",
            },
            "title": {"type": "string", "description": "Hujjat sarlavhasi."},
            "text": {
                "type": "string",
                "description": "docx, pdf va txt uchun mazmun (markdown).",
            },
            "rows": {
                "type": "array",
                "description": "xlsx va csv uchun qatorlar; birinchisi sarlavha.",
                "items": {"type": "array", "items": {"type": ["string", "number", "null"]}},
            },
            "sheets": {
                "type": "array",
                "description": "Excel'da bir nechta varaq kerak bo'lganda.",
                "items": {
                    "type": "object",
                    "properties": {
                        "name": {"type": "string"},
                        "rows": {
                            "type": "array",
                            "items": {"type": "array", "items": {"type": ["string", "number", "null"]}},
                        },
                    },
                    "required": ["name", "rows"],
                },
            },
        },
        "required": ["kind"],
    },
}


class BadInput(Exception):
    """Model noto'g'ri ma'lumot berdi — unga tushuntirib qaytaramiz."""


def run(payload) -> tuple[str, Artifact]:
    """Qurol chaqiruvini bajaradi: (modelga javob, tayyor fayl)."""
    if not isinstance(payload, dict):
        raise BadInput("Qurolga obyekt berilishi kerak edi.")

    kind = str(payload.get("kind") or "").strip().lower()
    if kind not in KINDS:
        raise BadInput(f"`kind` shulardan biri bo'lsin: {', '.join(KINDS)}.")

    title = str(payload.get("title") or "").strip()
    sheets = _clean_sheets(payload)
    text = payload.get("text")

    if kind in ("xlsx", "csv") and not sheets:
        raise BadInput("Jadval uchun `rows` yoki `sheets` berilishi kerak.")
    if kind in ("docx", "pdf", "txt") and not (isinstance(text, str) and text.strip()):
        raise BadInput("Hujjat uchun `text` berilishi kerak.")

    title = title or (sheets[0][0] if sheets else pdf.title_of(text))
    name = _filename(payload.get("filename") or title, kind)

    if kind == "xlsx":
        data = _xlsx(sheets)
    elif kind == "csv":
        data = _csv(sheets[0][1])
    elif kind == "docx":
        data = _docx(text, title)
    elif kind == "pdf":
        data = pdf.build(text, title)
    else:
        data = text.encode("utf-8")

    note = f"«{name}» tayyorlandi va foydalanuvchiga yuborildi."
    return note, Artifact(filename=name, data=data, kind=kind)


def _clean_sheets(payload: dict) -> list[tuple[str, list[list]]]:
    """`rows` yoki `sheets` ni bir xil ko'rinishga keltiradi."""
    sheets: list[tuple[str, list[list]]] = []

    raw_sheets = payload.get("sheets")
    if isinstance(raw_sheets, list) and raw_sheets:
        for index, sheet in enumerate(raw_sheets, start=1):
            if not isinstance(sheet, dict):
                continue
            rows = _clean_rows(sheet.get("rows"))
            if rows:
                sheets.append((str(sheet.get("name") or f"Varaq {index}")[:31], rows))
        return sheets

    rows = _clean_rows(payload.get("rows"))
    if rows:
        sheets.append((str(payload.get("title") or "Jadval")[:31], rows))
    return sheets


def _clean_rows(raw) -> list[list]:
    if not isinstance(raw, list):
        return []
    rows: list[list] = []
    for row in raw[:MAX_ROWS]:
        if isinstance(row, list):
            rows.append(["" if cell is None else cell for cell in row])
        elif row is not None:
            # Model bitta qatorni qator emas, matn qilib yuborishi mumkin.
            rows.append([row])
    return rows


def _filename(base: str, kind: str) -> str:
    slug = pdf.filename_of(str(base))[:-4]  # `.pdf` ni olib tashlaymiz
    return (slug or "fayl") + EXTENSIONS[kind]


def _csv(rows: list[list]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerows(rows)
    # Excel CSV'ni UTF-8 BOM bilan to'g'ri ochadi.
    return buffer.getvalue().encode("utf-8-sig")


def _xlsx(sheets: list[tuple[str, list[list]]]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    book.remove(book.active)

    for name, rows in sheets:
        sheet = book.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
        if rows:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            sheet.freeze_panes = "A2"
        _fit_columns(sheet, rows, get_column_letter)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _fit_columns(sheet, rows: list[list], letter) -> None:
    """Ustun kengligi — eng uzun qiymatga qarab, lekin cheklangan."""
    widths: dict[int, int] = {}
    for row in rows:
        for index, cell in enumerate(row, start=1):
            widths[index] = max(widths.get(index, 0), len(str(cell)))
    for index, width in widths.items():
        sheet.column_dimensions[letter(index)].width = min(max(width + 2, 8), 50)


def _docx(text: str, title: str) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(title, level=1)

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        heading = _HEADING.match(line)
        if heading:
            document.add_heading(_plain(heading.group(2)), level=2)
            continue
        bullet = _BULLET.match(line)
        if bullet:
            document.add_paragraph(_plain(bullet.group(1)), style="List Bullet")
            continue
        document.add_paragraph(_plain(line))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _plain(line: str) -> str:
    return _BOLD.sub(r"\1", line)
