"""Telegram'dan kelgan rasm va fayllarni Claude tushunadigan ko'rinishga o'girish.

Claude to'g'ridan-to'g'ri **rasm** va **PDF** ni o'qiy oladi — ular xabarga
maxsus blok sifatida qo'shiladi. Word va Excel esa unday emas: ulardan matnni
o'zimiz ajratib olib, oddiy matn sifatida yuboramiz.

Modul aiogram'ga bog'liq emas — faylni yuklab olish `main.py` ning ishi, bu
yerda faqat baytlar bilan ishlanadi. Shu sababli testlash oson.
"""
from __future__ import annotations

import base64
import io
import mimetypes

IMAGE_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}
PDF_TYPE = "application/pdf"
TEXT_TYPES = {
    "text/plain", "text/csv", "text/markdown", "text/html",
    "application/json", "application/xml", "text/xml",
}
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Ajratib olingan matnning chegarasi. Kattaroq fayl kontekstni ham, hisobni ham
# bo'shatadi — shuning uchun kesib, foydalanuvchini ogohlantiramiz.
MAX_TEXT_CHARS = 40_000
MAX_SHEET_ROWS = 300


def kind(mime: str | None, filename: str | None = None) -> str:
    """Fayl qanday ishlanishini aytadi: image, pdf, text, docx, xlsx, yo'q."""
    mime = (mime or "").split(";")[0].strip().lower()
    if not mime and filename:
        mime = (mimetypes.guess_type(filename)[0] or "").lower()

    if mime in IMAGE_TYPES:
        return "image"
    if mime == PDF_TYPE:
        return "pdf"
    if mime in TEXT_TYPES:
        return "text"
    if mime == DOCX_TYPE:
        return "docx"
    if mime == XLSX_TYPE:
        return "xlsx"

    # Mime kelmasa ham kengaytma bo'yicha urinib ko'ramiz.
    name = (filename or "").lower()
    if name.endswith((".txt", ".csv", ".md", ".json", ".log", ".py", ".sql")):
        return "text"
    if name.endswith(".docx"):
        return "docx"
    if name.endswith(".xlsx"):
        return "xlsx"
    return "unsupported"


def image_block(data: bytes, mime: str) -> dict:
    return {
        "type": "image",
        "source": {
            "type": "base64",
            "media_type": mime if mime in IMAGE_TYPES else "image/jpeg",
            "data": base64.standard_b64encode(data).decode("ascii"),
        },
    }


def pdf_block(data: bytes) -> dict:
    return {
        "type": "document",
        "source": {
            "type": "base64",
            "media_type": PDF_TYPE,
            "data": base64.standard_b64encode(data).decode("ascii"),
        },
    }


def text_block(body: str, filename: str) -> dict:
    """Matnli faylni oddiy matn bloki qilib beradi — nomi bilan birga."""
    body, cut = _clip(body)
    note = "\n\n[…fayl uzun bo'lgani uchun qisqartirildi]" if cut else ""
    return {
        "type": "text",
        "text": f"«{filename}» faylining mazmuni:\n\n{body}{note}",
    }


def extract_docx(data: bytes) -> str:
    """Word hujjatidan matn: xatboshi va jadvallar."""
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_xlsx(data: bytes) -> str:
    """Excel'dan matn: har varaq alohida, qatorlar tab bilan."""
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in book.worksheets:
            parts.append(f"## Varaq: {sheet.title}")
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_SHEET_ROWS:
                    parts.append(f"[…{sheet.title} varag'ining qolgan qatorlari tashlandi]")
                    break
                cells = ["" if value is None else str(value) for value in row]
                if any(cell.strip() for cell in cells):
                    parts.append("\t".join(cells))
    finally:
        book.close()
    return "\n".join(parts)


def _clip(body: str) -> tuple[str, bool]:
    body = body.strip()
    if len(body) <= MAX_TEXT_CHARS:
        return body, False
    return body[:MAX_TEXT_CHARS], True
