"""Claude Telegram boti: formatlash, tarix va API qatlami testlari.

Telegram'ga ham, Anthropic'ga ham chiqmaydi — hammasi lokal.
"""
import pytest

from claude_bot import formatting
from claude_bot.session import ChatHistory, RateLimiter


# --- formatting ------------------------------------------------------------

def test_html_escapes_dangerous_characters():
    html = formatting.to_html("5 < 7 & <script>alert(1)</script>")
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert "5 &lt; 7 &amp; " in html


def test_markdown_becomes_telegram_html():
    html = formatting.to_html("## Sarlavha\n**qalin** va *egik*\n- birinchi")
    assert "<b>Sarlavha</b>" in html
    assert "<b>qalin</b>" in html
    assert "<i>egik</i>" in html
    assert "• birinchi" in html


def test_code_block_keeps_content_and_drops_language_tag():
    html = formatting.to_html("Mana:\n```python\nprint('a' < 'b')\n```")
    assert "<pre><code>print(&#x27;a&#x27; &lt; &#x27;b&#x27;)</code></pre>" in html
    assert "python" not in html


def test_stars_inside_code_are_not_formatting():
    html = formatting.to_html("`a ** b` matn")
    assert "<code>a ** b</code>" in html
    assert "<b>" not in html


def test_links_become_anchors():
    html = formatting.to_html("[sayt](https://example.com) ochiladi")
    assert '<a href="https://example.com">sayt</a>' in html


def test_short_answer_is_one_chunk():
    assert formatting.split_markdown("qisqa javob") == ["qisqa javob"]


def test_long_answer_is_split_under_the_limit():
    text = "\n".join(f"{index}-satr" for index in range(2000))
    chunks = formatting.split_markdown(text, limit=500)
    assert len(chunks) > 1
    assert all(len(chunk) <= 500 for chunk in chunks)
    # Hech bir satr yo'qolmasligi kerak.
    assert "1999-satr" in "\n".join(chunks)


def test_split_closes_and_reopens_a_code_fence():
    body = "\n".join(f"line_{index}" for index in range(200))
    chunks = formatting.split_markdown(f"```python\n{body}\n```", limit=400)
    assert len(chunks) > 1
    for chunk in chunks:
        # Har bir bo'lakda ``` juft bo'lsin — aks holda HTML buziladi.
        assert chunk.count("```") % 2 == 0
    for chunk in chunks:
        assert "<pre><code>" in formatting.to_html(chunk)


def test_very_long_single_line_is_cut():
    chunks = formatting.split_markdown("x" * 900, limit=300)
    assert [len(chunk) for chunk in chunks] == [300, 300, 300]


def test_render_produces_html_chunks_aligned_with_plain_chunks():
    text = "\n".join(f"{index}. qator" for index in range(500))
    assert len(formatting.render(text, limit=400)) == len(
        formatting.split_markdown(text, limit=400)
    )


# --- tarix va cheklov -------------------------------------------------------

def test_history_keeps_only_the_last_messages():
    history = ChatHistory(limit=4)
    for index in range(5):
        history.add(1, "user", f"savol {index}")
        history.add(1, "assistant", f"javob {index}")

    messages = history.get(1)
    assert len(messages) <= 4
    # API birinchi xabar `user` bo'lishini talab qiladi.
    assert messages[0]["role"] == "user"
    assert messages[-1]["content"] == "javob 4"


def test_history_is_per_chat_and_clearable():
    history = ChatHistory()
    history.add(1, "user", "birinchi chat")
    history.add(2, "user", "ikkinchi chat")
    history.clear(1)
    assert history.get(1) == []
    assert history.get(2)[0]["content"] == "ikkinchi chat"


def test_rate_limiter_blocks_then_recovers():
    limiter = RateLimiter(per_minute=2)
    assert limiter.allow(7, now=0.0)
    assert limiter.allow(7, now=1.0)
    assert not limiter.allow(7, now=2.0)
    assert limiter.allow(8, now=2.0)  # boshqa foydalanuvchiga ta'sir qilmaydi
    assert limiter.allow(7, now=100.0)  # oyna o'tdi


def test_rate_limiter_can_be_disabled():
    limiter = RateLimiter(per_minute=0)
    assert all(limiter.allow(1, now=float(step)) for step in range(50))


# --- Claude qatlami ---------------------------------------------------------

class _Delta:
    def __init__(self, text):
        self.type = "text_delta"
        self.text = text


class _Event:
    """Stream hodisasi: matn bo'lagi yoki qurol ishga tushishi."""

    def __init__(self, type_, delta=None, content_block=None):
        self.type = type_
        self.delta = delta
        self.content_block = content_block

    @staticmethod
    def text(piece):
        return _Event("content_block_delta", delta=_Delta(piece))

    @staticmethod
    def tool(name):
        block = type("Block", (), {"type": "server_tool_use", "name": name})()
        return _Event("content_block_start", content_block=block)


class _FakeStream:
    """anthropic'ning stream helper'iga o'xshab qiladigan soxta obyekt."""

    def __init__(self, pieces, final, events=None):
        # `pieces` — qulaylik uchun: matn bo'laklari hodisaga aylantiriladi.
        self._events = events if events is not None else [
            _Event.text(piece) for piece in pieces
        ]
        self._final = final

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc_info):
        return False

    def __aiter__(self):
        async def generator():
            for event in self._events:
                yield event
        return generator()

    async def get_final_message(self):
        return self._final


class _Block:
    def __init__(self, text):
        self.type = "text"
        self.text = text


class _Usage:
    input_tokens = 11
    output_tokens = 22


class _ToolUse:
    def __init__(self, tool_id, name, payload):
        self.type = "tool_use"
        self.id = tool_id
        self.name = name
        self.input = payload


class _Message:
    def __init__(self, text, stop_reason="end_turn", stop_details=None, content=None):
        self.content = content if content is not None else [_Block(text)]
        self.stop_reason = stop_reason
        self.stop_details = stop_details
        self.usage = _Usage()
        self.model = "claude-opus-5"


def _client(monkeypatch, stream, captured=None, **options):
    pytest.importorskip("anthropic")
    from claude_bot.claude import ClaudeClient

    client = ClaudeClient(
        "sk-test",
        model="claude-opus-5",
        system="test",
        max_tokens=100,
        effort="low",
        use_fallbacks=False,
        **options,
    )

    def fake_stream(**kwargs):
        if captured is not None:
            captured.update(kwargs)
        return stream

    monkeypatch.setattr(client._client.messages, "stream", fake_stream, raising=False)
    return client


def test_ask_collects_streamed_text_and_reports_progress(monkeypatch):
    import asyncio

    stream = _FakeStream(["Sa", "lom", "!"], _Message("Salom!"))
    client = _client(monkeypatch, stream)

    seen = []

    async def on_progress(partial):
        seen.append(partial)

    async def run():
        return await client.ask([{"role": "user", "content": "hi"}], on_progress)

    reply = asyncio.run(run())
    assert reply.text == "Salom!"
    assert seen == ["Sa", "Salom", "Salom!"]
    assert (reply.input_tokens, reply.output_tokens) == (11, 22)
    assert not reply.truncated


def test_ask_surfaces_a_refusal(monkeypatch):
    import asyncio

    class _Details:
        category = "cyber"
        explanation = "not allowed"

    stream = _FakeStream([], _Message("", stop_reason="refusal", stop_details=_Details()))
    client = _client(monkeypatch, stream)

    reply = asyncio.run(client.ask([{"role": "user", "content": "hi"}]))
    assert reply.stop_reason == "refusal"
    assert reply.refusal == "not allowed"


def test_truncated_answer_is_flagged(monkeypatch):
    import asyncio

    stream = _FakeStream(["uzun"], _Message("uzun", stop_reason="max_tokens"))
    client = _client(monkeypatch, stream)

    reply = asyncio.run(client.ask([{"role": "user", "content": "hi"}]))
    assert reply.truncated


# --- handler mantiqi --------------------------------------------------------

class _StubChat:
    def __init__(self, chat_id):
        self.id = chat_id


class _StubUser:
    def __init__(self, user_id):
        self.id = user_id


class _StubMessage:
    """Telegram xabarining eng kerakli qismi — aiogram'siz."""

    def __init__(self, text="", chat_id=1, user_id=1, outbox=None, edits=None,
                 photo=None, document=None, caption=None):
        self.text = text
        self.photo = photo
        self.document = document
        self.caption = caption
        self.chat = _StubChat(chat_id)
        self.from_user = _StubUser(user_id)
        # Ro'yxatlar «chat» bo'ylab umumiy: placeholder alohida obyekt bo'lsa ham
        # uning tahrirlari shu yerda ko'rinadi.
        self.outbox = outbox if outbox is not None else []
        self.edits = edits if edits is not None else []

    async def answer(self, text, **kwargs):
        self.outbox.append(text)
        return _StubMessage(
            text, self.chat.id, self.from_user.id, self.outbox, self.edits
        )

    async def edit_text(self, text, **kwargs):
        self.edits.append(text)
        self.text = text

    async def answer_document(self, document, **kwargs):
        self.outbox.append(document)
        return self


class _StubBot:
    def __init__(self, payload=b""):
        self.actions = []
        self.payload = payload

    async def send_chat_action(self, chat_id, action):
        self.actions.append((chat_id, action))

    async def get_file(self, file_id):
        return type("Info", (), {"file_path": f"files/{file_id}"})()

    async def download_file(self, path):
        import io
        return io.BytesIO(self.payload)


class _StubClaude:
    def __init__(self, reply):
        self.reply = reply
        self.seen = None

    async def ask(self, messages, on_progress=None, on_status=None, on_tool=None):
        self.seen = list(messages)
        if on_status is not None:
            await on_status("web_search")
        if on_progress is not None:
            await on_progress(self.reply.text)
        return self.reply


def _run_question(monkeypatch, reply, *, text="Salom", history=None, limiter=None):
    import asyncio

    pytest.importorskip("aiogram")
    from claude_bot import main
    from claude_bot.session import UsageTracker

    message = _StubMessage(text)
    claude = _StubClaude(reply)
    history = history or ChatHistory()
    asyncio.run(
        main.on_question(
            message,
            bot=_StubBot(),
            claude=claude,
            history=history,
            limiter=limiter or RateLimiter(per_minute=0),
            usage=UsageTracker(),
        )
    )
    return message, claude, history


def _reply(text, **kwargs):
    pytest.importorskip("anthropic")
    from claude_bot.claude import Reply

    return Reply(text=text, **kwargs)


def test_question_reaches_claude_and_answer_reaches_the_chat(monkeypatch):
    message, claude, history = _run_question(
        monkeypatch, _reply("**Javob** bu yerda"), text="Savolim"
    )

    assert claude.seen == [{"role": "user", "content": "Savolim"}]
    # Birinchi xabar «o'ylayapman», keyin u javob bilan tahrirlanadi.
    assert message.outbox and message.outbox[0]
    assert "<b>Javob</b> bu yerda" in message.edits[-1]
    assert [item["content"] for item in history.get(1)] == [
        "Savolim",
        "**Javob** bu yerda",
    ]


def test_long_answer_is_delivered_in_several_messages(monkeypatch):
    long_answer = "\n".join(f"{index}-qator" for index in range(2000))
    message, _, _ = _run_question(monkeypatch, _reply(long_answer))

    expected = len(formatting.split_markdown(long_answer))
    assert expected > 1
    # Bittasi «o'ylayapman» xabari, qolganlari — javob bo'laklari.
    assert len(message.outbox) == 1 + (expected - 1)


def test_refusal_is_explained_and_not_remembered(monkeypatch):
    message, _, history = _run_question(
        monkeypatch, _reply("", stop_reason="refusal", refusal="cyber")
    )
    assert "bosh tortdi" in message.edits[-1]
    assert history.get(1) == []


def test_rate_limited_user_is_not_sent_to_claude(monkeypatch):
    limiter = RateLimiter(per_minute=1)
    limiter.allow(1)
    message, claude, _ = _run_question(
        monkeypatch, _reply("javob"), limiter=limiter
    )
    assert claude.seen is None
    assert "sekinroq" in message.outbox[0]


def test_overlong_question_is_rejected(monkeypatch):
    from claude_bot import config

    message, claude, _ = _run_question(
        monkeypatch, _reply("javob"), text="x" * (config.MAX_QUESTION_CHARS + 1)
    )
    assert claude.seen is None
    assert "uzun" in message.outbox[0]


def test_request_carries_thinking_and_effort(monkeypatch):
    import asyncio

    captured = {}
    client = _client(monkeypatch, _FakeStream(["ok"], _Message("ok")), captured)
    asyncio.run(client.ask([{"role": "user", "content": "hi"}]))

    assert captured["thinking"] == {"type": "adaptive"}
    assert captured["output_config"] == {"effort": "low"}
    assert captured["messages"] == [{"role": "user", "content": "hi"}]


def test_thinking_can_be_turned_off_for_older_models(monkeypatch):
    import asyncio

    captured = {}
    client = _client(
        monkeypatch, _FakeStream(["ok"], _Message("ok")), captured, use_thinking=False
    )
    asyncio.run(client.ask([{"role": "user", "content": "hi"}]))

    assert "thinking" not in captured


# --- rasm va fayl -----------------------------------------------------------

def test_kind_recognises_what_we_can_read():
    from claude_bot import media

    assert media.kind("image/png", "a.png") == "image"
    assert media.kind("application/pdf", "hisobot.pdf") == "pdf"
    assert media.kind("text/csv", "data.csv") == "text"
    assert media.kind(media.DOCX_TYPE, "shartnoma.docx") == "docx"
    assert media.kind(media.XLSX_TYPE, "byudjet.xlsx") == "xlsx"
    assert media.kind("video/mp4", "klip.mp4") == "unsupported"


def test_kind_falls_back_to_the_file_name():
    from claude_bot import media

    # Telegram mime bermasligi mumkin — kengaytma bo'yicha topamiz.
    assert media.kind(None, "hisobot.docx") == "docx"
    assert media.kind("", "eslatma.txt") == "text"
    assert media.kind("application/octet-stream", "jadval.xlsx") == "xlsx"


def test_image_block_is_base64_with_media_type():
    import base64

    from claude_bot import media

    block = media.image_block(b"\x89PNG-fake", "image/png")
    assert block["type"] == "image"
    assert block["source"]["media_type"] == "image/png"
    assert base64.standard_b64decode(block["source"]["data"]) == b"\x89PNG-fake"


def test_unknown_image_type_falls_back_to_jpeg():
    from claude_bot import media

    block = media.image_block(b"data", "image/heic")
    assert block["source"]["media_type"] == "image/jpeg"


def test_pdf_becomes_a_document_block():
    from claude_bot import media

    block = media.pdf_block(b"%PDF-1.4")
    assert block["type"] == "document"
    assert block["source"]["media_type"] == media.PDF_TYPE


def test_long_text_file_is_clipped_and_says_so():
    from claude_bot import media

    block = media.text_block("x" * (media.MAX_TEXT_CHARS + 500), "katta.txt")
    assert "katta.txt" in block["text"]
    assert "qisqartirildi" in block["text"]
    assert len(block["text"]) < media.MAX_TEXT_CHARS + 500


def test_word_document_text_is_extracted():
    docx = pytest.importorskip("docx")
    import io

    from claude_bot import media

    document = docx.Document()
    document.add_paragraph("Shartnoma raqami 42")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Summa"
    table.rows[0].cells[1].text = "1000"
    buffer = io.BytesIO()
    document.save(buffer)

    body = media.extract_docx(buffer.getvalue())
    assert "Shartnoma raqami 42" in body
    assert "Summa | 1000" in body


def test_excel_sheet_text_is_extracted():
    openpyxl = pytest.importorskip("openpyxl")
    import io

    from claude_bot import media

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Byudjet"
    sheet.append(["Oy", "Xarajat"])
    sheet.append(["Yanvar", 1500])
    buffer = io.BytesIO()
    book.save(buffer)

    body = media.extract_xlsx(buffer.getvalue())
    assert "## Varaq: Byudjet" in body
    assert "Yanvar\t1500" in body


def test_only_the_newest_file_stays_in_context():
    history = ChatHistory()
    first = [{"type": "image", "source": {"data": "aaa"}}]
    second = [{"type": "image", "source": {"data": "bbb"}}]

    history.add(1, "user", first, marker="[rasm-1 yuborildi]")
    history.add(1, "assistant", "birinchi javob")
    history.add(1, "user", second, marker="[rasm-2 yuborildi]")

    messages = history.get(1)
    # Eski rasm o'z belgisiga almashdi, yangisi bloklarcha turibdi.
    assert messages[0]["content"] == "[rasm-1 yuborildi]"
    assert messages[2]["content"] == second


def test_plain_text_history_is_untouched_by_markers():
    history = ChatHistory()
    history.add(1, "user", "oddiy savol")
    history.add(1, "assistant", "javob")
    history.add(1, "user", "yana savol")

    assert [item["content"] for item in history.get(1)] == [
        "oddiy savol", "javob", "yana savol",
    ]


# --- jadval bo'yicha ishlash ------------------------------------------------

def _task(hour=8, minute=0, days="daily", last_run=None, prompt="yangiliklar"):
    from claude_bot.tasks import Task

    return Task(1, 100, prompt, hour, minute, days, last_run)


def _at(year=2026, month=9, day=16, hour=8, minute=0):
    from datetime import datetime, timedelta, timezone

    return datetime(year, month, day, hour, minute, tzinfo=timezone(timedelta(hours=5)))


def test_task_runs_at_its_time_and_not_before():
    from claude_bot.tasks import is_due

    assert is_due(_task(hour=8), _at(hour=8, minute=0))
    assert is_due(_task(hour=8), _at(hour=8, minute=30))
    assert not is_due(_task(hour=8), _at(hour=7, minute=59))


def test_task_runs_once_a_day():
    from claude_bot.tasks import is_due

    # 2026-09-16 — chorshanba.
    already = _task(hour=8, last_run="2026-09-16")
    assert not is_due(already, _at(hour=9))
    assert is_due(already, _at(day=17, hour=9))


def test_late_wake_still_runs_but_not_at_midnight():
    from claude_bot.tasks import is_due

    task = _task(hour=8)
    # Servis uxlab qolib, 12:00 da uyg'ondi — ertalabki vazifa baribir kerak.
    assert is_due(task, _at(hour=12), grace_hours=6)
    # Kechqurun esa ertalabki xulosa keraksiz.
    assert not is_due(task, _at(hour=22), grace_hours=6)


def test_weekday_tasks_respect_the_day():
    from claude_bot.tasks import is_due

    # 16-sentabr 2026 — chorshanba (weekday=2), 19-si — shanba (weekday=5).
    workday = _task(days="workdays")
    assert is_due(workday, _at(day=16, hour=9))
    assert not is_due(workday, _at(day=19, hour=9))

    monday_only = _task(days="0")
    assert not is_due(monday_only, _at(day=16, hour=9))
    assert is_due(monday_only, _at(day=21, hour=9))  # 21-si — dushanba


def test_days_are_understood_in_uzbek_and_english():
    from claude_bot import tasks

    assert tasks.parse_days(None) == tasks.DAILY
    assert tasks.parse_days("har kuni") == tasks.DAILY
    assert tasks.parse_days("ish kunlari") == tasks.WORKDAYS
    assert tasks.parse_days("dushanba, juma") == "0,4"
    assert tasks.parse_days("mon,wed") == "0,2"
    assert tasks.parse_days("allaqanday matn") == tasks.DAILY


def test_time_is_parsed_or_refused_clearly():
    from claude_bot import tasks

    assert tasks.parse_time("08:00") == (8, 0)
    assert tasks.parse_time("8") == (8, 0)
    assert tasks.parse_time("9.30") == (9, 30)
    with pytest.raises(ValueError):
        tasks.parse_time("ertalab")
    with pytest.raises(ValueError):
        tasks.parse_time("25:00")


def test_task_description_reads_like_a_sentence():
    assert _task(hour=8, minute=5).when == "har kuni 08:05"
    assert _task(days="workdays").when == "ish kunlari 08:00"
    assert _task(days="0,4").when == "dushanba, juma 08:00"


def test_tasks_survive_a_restart(tmp_path):
    from claude_bot.tasks import TaskStore

    store = TaskStore(tmp_path / "tasks.db")
    store.add(100, "yangiliklarni yubor", 8, 0, "daily")
    store.close()

    # Bot qayta ishga tushdi — vazifa joyida bo'lishi kerak.
    again = TaskStore(tmp_path / "tasks.db")
    items = again.for_chat(100)
    assert len(items) == 1
    assert items[0].prompt == "yangiliklarni yubor"


def test_tasks_are_per_chat_and_removable(tmp_path):
    from claude_bot.tasks import TaskStore

    store = TaskStore(tmp_path / "tasks.db")
    mine = store.add(100, "meniki", 8, 0, "daily")
    store.add(200, "boshqasi", 9, 0, "daily")

    # Boshqa chat meni vazifamni o'chira olmasin.
    assert not store.remove(200, mine.id)
    assert store.remove(100, mine.id)
    assert store.for_chat(100) == []
    assert len(store.for_chat(200)) == 1


def test_a_chat_cannot_pile_up_tasks(tmp_path):
    from claude_bot import tasks

    store = tasks.TaskStore(tmp_path / "tasks.db")
    for index in range(tasks.MAX_TASKS_PER_CHAT):
        store.add(100, f"vazifa {index}", 8, 0, "daily")

    with pytest.raises(ValueError):
        store.add(100, "ortiqcha", 9, 0, "daily")


def test_due_task_is_run_and_marked(tmp_path):
    import asyncio

    pytest.importorskip("aiogram")
    from claude_bot import scheduler, tasks

    store = tasks.TaskStore(tmp_path / "tasks.db")
    store.add(100, "kursni yubor", 0, 0, "daily")

    sent = []

    class _Bot:
        async def send_message(self, chat_id, text, **kwargs):
            sent.append((chat_id, text))

    class _Claude:
        async def ask(self, messages, *args, **kwargs):
            return _reply("Kurs 11 774 so'm")

    count = asyncio.run(scheduler.run_due(
        _Bot(), _Claude(), store, tz_offset=5, grace_hours=24,
    ))

    assert count == 1
    assert sent[0][0] == 100
    assert any("11 774" in text for _, text in sent)
    # Ikkinchi marta bajarilmasin.
    assert asyncio.run(scheduler.run_due(
        _Bot(), _Claude(), store, tz_offset=5, grace_hours=24,
    )) == 0


def test_failed_task_still_reaches_the_user(tmp_path):
    import asyncio

    pytest.importorskip("aiogram")
    from claude_bot import scheduler, tasks

    store = tasks.TaskStore(tmp_path / "tasks.db")
    store.add(100, "nimadir", 0, 0, "daily")

    sent = []

    class _Bot:
        async def send_message(self, chat_id, text, **kwargs):
            sent.append(text)

    class _Claude:
        async def ask(self, messages, *args, **kwargs):
            raise RuntimeError("tarmoq yiqildi")

    asyncio.run(scheduler.run_due(_Bot(), _Claude(), store, tz_offset=5, grace_hours=24))

    # Jim qolgandan ko'ra, xatolikni aytgan yaxshi.
    assert any("❌" in text for text in sent)


def test_schedule_tool_adds_lists_and_cancels(tmp_path):
    pytest.importorskip("aiogram")
    from claude_bot import main, tasks

    store = tasks.TaskStore(tmp_path / "tasks.db")

    added = main._run_task_tool(
        "schedule_task",
        {"time": "08:00", "days": "ish kunlari", "prompt": "Yangiliklarni xulosa qil"},
        store, 100,
    )
    assert "ish kunlari 08:00" in added

    listed = main._run_task_tool("list_tasks", {}, store, 100)
    assert "Yangiliklarni xulosa qil" in listed

    task_id = store.for_chat(100)[0].id
    assert "o'chirildi" in main._run_task_tool("cancel_task", {"id": task_id}, store, 100)
    assert store.for_chat(100) == []


def test_schedule_tool_explains_bad_input(tmp_path):
    pytest.importorskip("aiogram")
    from claude_bot import files, main, tasks

    store = tasks.TaskStore(tmp_path / "tasks.db")

    with pytest.raises(files.BadInput):
        main._run_task_tool("schedule_task", {"time": "ertalab", "prompt": "x"}, store, 100)
    with pytest.raises(files.BadInput):
        main._run_task_tool("schedule_task", {"time": "08:00", "prompt": " "}, store, 100)


def test_tick_endpoint_needs_the_key():
    import asyncio

    pytest.importorskip("aiogram")
    from claude_bot import webhook

    calls = []

    async def on_tick():
        calls.append(1)
        return 2

    handler = webhook._make_tick_handler(on_tick, "maxfiy")

    class _Request:
        def __init__(self, key):
            self.query = {"key": key} if key else {}

    denied = asyncio.run(handler(_Request("boshqa")))
    assert denied.status == 403
    assert calls == []

    allowed = asyncio.run(handler(_Request("maxfiy")))
    assert "ran 2" in allowed.text
    assert calls == [1]


# --- fayl yasash quroli -----------------------------------------------------

def test_excel_is_built_from_rows():
    pytest.importorskip("openpyxl")
    import io

    import openpyxl

    from claude_bot import files

    note, artifact = files.run({
        "kind": "xlsx",
        "filename": "qarindoshlar",
        "rows": [["Ism", "Summa"], ["Akram", 300000], ["Elnur", 800000]],
    })

    assert artifact.filename == "qarindoshlar.xlsx"
    assert "qarindoshlar.xlsx" in note
    sheet = openpyxl.load_workbook(io.BytesIO(artifact.data)).worksheets[0]
    assert [cell.value for cell in sheet[1]] == ["Ism", "Summa"]
    assert sheet["B3"].value == 800000
    assert sheet[1][0].font.bold  # sarlavha ajratilgan bo'lsin


def test_excel_can_hold_several_sheets():
    pytest.importorskip("openpyxl")
    import io

    import openpyxl

    from claude_bot import files

    _, artifact = files.run({
        "kind": "xlsx",
        "sheets": [
            {"name": "Yanvar", "rows": [["a", 1]]},
            {"name": "Fevral", "rows": [["b", 2]]},
        ],
    })

    book = openpyxl.load_workbook(io.BytesIO(artifact.data))
    assert book.sheetnames == ["Yanvar", "Fevral"]


def test_word_document_is_built_from_text():
    docx = pytest.importorskip("docx")
    import io

    from claude_bot import files

    _, artifact = files.run({
        "kind": "docx",
        "title": "Hisobot",
        "text": "## Bo'lim\n\n- Birinchi band\n\nOddiy xatboshi.",
    })

    assert artifact.filename == "hisobot.docx"
    document = docx.Document(io.BytesIO(artifact.data))
    texts_in_doc = [p.text for p in document.paragraphs]
    assert "Hisobot" in texts_in_doc
    assert "Birinchi band" in texts_in_doc


def test_csv_is_utf8_with_bom_so_excel_opens_it():
    from claude_bot import files

    _, artifact = files.run({"kind": "csv", "rows": [["Ism", "Summa"], ["Aziz", 10]]})

    assert artifact.filename.endswith(".csv")
    assert artifact.data.startswith(b"\xef\xbb\xbf")
    assert b"Ism,Summa" in artifact.data


def test_pdf_can_be_made_through_the_tool():
    pytest.importorskip("reportlab")
    from claude_bot import files

    _, artifact = files.run({"kind": "pdf", "title": "Bayon", "text": "Matn"})
    assert artifact.data.startswith(b"%PDF")


def test_tool_refuses_what_it_cannot_build():
    from claude_bot import files

    with pytest.raises(files.BadInput):
        files.run({"kind": "mp3", "text": "salom"})
    with pytest.raises(files.BadInput):
        files.run({"kind": "xlsx"})           # jadvalsiz jadval
    with pytest.raises(files.BadInput):
        files.run({"kind": "docx", "text": " "})  # bo'sh hujjat
    with pytest.raises(files.BadInput):
        files.run("umuman boshqa narsa")


def test_stray_row_shapes_are_tolerated():
    from claude_bot import files

    # Model qatorni matn qilib yuborishi mumkin — yiqilmasin.
    _, artifact = files.run({"kind": "csv", "rows": ["birinchi", ["ikkinchi", 2], None]})
    assert b"birinchi" in artifact.data
    assert b"ikkinchi,2" in artifact.data


def test_tool_call_reaches_the_user_as_a_document(monkeypatch):
    import asyncio

    pytest.importorskip("aiogram")
    pytest.importorskip("openpyxl")
    from claude_bot import main
    from claude_bot.session import UsageTracker

    class _FileClaude:
        """Birinchi navbatda qurolni chaqiradi, keyin javob beradi."""

        def __init__(self):
            self.reply = _reply("Excel tayyor", input_tokens=10, output_tokens=5)

        async def ask(self, messages, on_progress=None, on_status=None, on_tool=None):
            await on_tool("create_file", {
                "kind": "xlsx", "filename": "royxat",
                "rows": [["Ism", "Summa"], ["Akram", 300000]],
            })
            return self.reply

    message = _StubMessage("ro'yxatni excelga sol")
    asyncio.run(
        main.on_question(
            message,
            bot=_StubBot(),
            claude=_FileClaude(),
            history=ChatHistory(),
            limiter=RateLimiter(per_minute=0),
            usage=UsageTracker(),
        )
    )

    sent = [item for item in message.outbox if hasattr(item, "filename")]
    assert len(sent) == 1
    assert sent[0].filename == "royxat.xlsx"


# --- xarajat ----------------------------------------------------------------

def test_cost_follows_the_price_list():
    from claude_bot import pricing

    # 1M kirish + 1M chiqish = $5 + $25 Opus uchun.
    assert pricing.cost("claude-opus-5", 1_000_000, 1_000_000) == pytest.approx(30.0)
    assert pricing.cost("claude-sonnet-5", 1_000_000, 0) == pytest.approx(2.0)
    assert pricing.cost("claude-haiku-4-5", 0, 1_000_000) == pytest.approx(5.0)


def test_unknown_model_is_priced_high_rather_than_low():
    from claude_bot import pricing

    # Taxmin kam chiqqandan ko'ra ko'p chiqqani xavfsiz.
    assert pricing.rates("qandaydir-yangi-model") == pricing.FALLBACK


def test_money_reads_naturally():
    from claude_bot import pricing

    assert pricing.money(0) == "0"
    assert pricing.money(0.004) == "1 sentdan kam"
    assert pricing.money(0.03) == "~3 sent"
    assert pricing.money(1.5) == "~$1.50"


def test_usage_separates_the_last_day_from_the_total():
    from claude_bot.session import UsageTracker

    usage = UsageTracker()
    now = 1_000_000.0
    usage.record(1, 0.05, 1000, 500, now=now - 48 * 3600)  # ikki kun oldin
    usage.record(1, 0.02, 400, 200, now=now - 3600)        # bir soat oldin

    summary = usage.summary(1, now=now)
    assert summary["calls"] == 2
    assert summary["usd"] == pytest.approx(0.07)
    assert summary["recent_calls"] == 1
    assert summary["recent_usd"] == pytest.approx(0.02)


def test_usage_is_per_chat():
    from claude_bot.session import UsageTracker

    usage = UsageTracker()
    usage.record(1, 0.10, 100, 100)
    usage.record(2, 0.01, 10, 10)

    assert usage.summary(1)["usd"] == pytest.approx(0.10)
    assert usage.summary(2)["usd"] == pytest.approx(0.01)
    assert usage.summary(3)["calls"] == 0


def test_answer_carries_the_price_note(monkeypatch):
    from claude_bot import config

    monkeypatch.setattr(config, "SHOW_COST", True)
    reply = _reply("Javob", input_tokens=5000, output_tokens=2000)
    message, _, _ = _run_question(monkeypatch, reply)

    # Narx oxirgi bo'lakka qo'shiladi, javobning o'ziga tegmaydi.
    assert "Javob" in message.edits[-1]
    assert "sent" in message.edits[-1]


def test_price_note_can_be_turned_off(monkeypatch):
    from claude_bot import config

    monkeypatch.setattr(config, "SHOW_COST", False)
    reply = _reply("Javob", input_tokens=5000, output_tokens=2000)
    message, _, _ = _run_question(monkeypatch, reply)

    assert "sent" not in message.edits[-1]


def test_cost_command_reports_both_windows():
    import asyncio

    pytest.importorskip("aiogram")
    from claude_bot import main
    from claude_bot.session import UsageTracker

    usage = UsageTracker()
    usage.record(1, 0.12, 5000, 2000)

    message = _StubMessage()
    asyncio.run(main.cmd_cost(message, usage=usage))

    report = message.outbox[-1]
    assert "Sarf hisobi" in report
    assert "12 sent" in report
    assert "/new" in report  # tejash maslahati ham bo'lsin


# --- PDF --------------------------------------------------------------------

def test_pdf_is_a_real_pdf():
    pytest.importorskip("reportlab")
    from claude_bot import pdf

    data = pdf.build("## Hisobot\n\n- Birinchi\n- Ikkinchi\n\n**Jami: 100**")
    assert data.startswith(b"%PDF")
    assert len(data) > 800


def test_pdf_title_comes_from_the_first_heading():
    pytest.importorskip("reportlab")
    from claude_bot import pdf

    assert pdf.title_of("## Qarindoshlar o'tirishi\n\nmatn") == "Qarindoshlar o'tirishi"
    assert pdf.title_of("**Oylik hisobot**\nmatn") == "Oylik hisobot"
    assert pdf.title_of("\n\n") == "Javob"


def test_pdf_filename_is_safe():
    pytest.importorskip("reportlab")
    from claude_bot import pdf

    assert pdf.filename_of("Qarindoshlar o'tirishi") == "qarindoshlar_o_tirishi.pdf"
    assert pdf.filename_of("!!!").endswith(".pdf")
    assert " " not in pdf.filename_of("uzun nom bilan fayl")


def test_pdf_still_builds_without_a_unicode_font(monkeypatch):
    pytest.importorskip("reportlab")
    from claude_bot import pdf

    # Hostingda DejaVu bo'lmasligi mumkin — o'shanda ham PDF chiqishi kerak.
    monkeypatch.setattr(pdf, "FONT_PATHS", ())
    data = pdf.build("Belgilar: oʻzbek — «qo'shtirnoq» va ustunlar")
    assert data.startswith(b"%PDF")


# --- qurollar ---------------------------------------------------------------

def test_web_mode_offers_search_and_fetch():
    from claude_bot import tools

    kinds = {tool["type"] for tool in tools.build("web", max_uses=3)}
    assert kinds == {tools.WEB_SEARCH, tools.WEB_FETCH}
    assert all(tool["max_uses"] == 3 for tool in tools.build("web", max_uses=3))


def test_code_mode_avoids_two_execution_environments():
    from claude_bot import tools

    kinds = {tool["type"] for tool in tools.build("code")}
    # Yangi qidiruv quroli ichida kod bajarish muhiti bor — kod rejimida
    # qidiruvning oddiy varianti olinadi, aks holda muhit ikkita bo'lib qoladi.
    assert tools.CODE_EXECUTION in kinds
    assert tools.WEB_SEARCH not in kinds
    assert tools.BASIC_WEB_SEARCH in kinds


def test_tools_can_be_turned_off():
    from claude_bot import tools

    assert tools.build("off") == []


def test_request_carries_tools_when_enabled(monkeypatch):
    import asyncio

    captured = {}
    client = _client(
        monkeypatch,
        _FakeStream(["ok"], _Message("ok")),
        captured,
        tool_mode="web",
    )
    asyncio.run(client.ask([{"role": "user", "content": "hi"}]))

    assert [tool["name"] for tool in captured["tools"]] == ["web_search", "web_fetch"]


def test_no_tools_key_when_mode_is_off(monkeypatch):
    import asyncio

    captured = {}
    client = _client(monkeypatch, _FakeStream(["ok"], _Message("ok")), captured)
    asyncio.run(client.ask([{"role": "user", "content": "hi"}]))

    assert "tools" not in captured


def test_search_that_runs_code_still_reads_as_searching():
    from claude_bot import texts

    # Qidiruv quroli natijalarni saralash uchun ichida kod bajaradi — bu
    # foydalanuvchiga «hisoblayapman» bo'lib ko'rinmasligi kerak.
    assert texts.tool_status("code_execution", "web") == texts.TOOL_STATUS["web_search"]
    # `code` rejimida esa hisoblash haqiqatan hisoblash.
    assert texts.tool_status("code_execution", "code") == texts.TOOL_STATUS["code_execution"]
    assert texts.tool_status("web_fetch", "web") == texts.TOOL_STATUS["web_fetch"]


def test_status_callback_fires_when_a_tool_starts(monkeypatch):
    import asyncio

    stream = _FakeStream(
        None,
        _Message("javob"),
        events=[_Event.tool("web_search"), _Event.text("javob")],
    )
    client = _client(monkeypatch, stream, tool_mode="web")

    seen = []

    async def on_status(name):
        seen.append(name)

    asyncio.run(
        client.ask([{"role": "user", "content": "hi"}], None, on_status)
    )
    assert seen == ["web_search"]


def _tool_client(monkeypatch, replies, sent, on_tool_calls):
    pytest.importorskip("anthropic")
    from claude_bot.claude import ClaudeClient

    client = ClaudeClient(
        "sk-test", model="claude-opus-5", system="test",
        max_tokens=100, effort="low", use_fallbacks=False,
    )

    def fake_stream(**kwargs):
        sent.append(kwargs["messages"])
        return replies[len(sent) - 1]

    monkeypatch.setattr(client._client.messages, "stream", fake_stream, raising=False)
    return client


def test_tool_call_is_executed_and_the_turn_continues(monkeypatch):
    import asyncio

    sent, calls = [], []
    replies = [
        _FakeStream([], _Message("", stop_reason="tool_use", content=[
            _ToolUse("t1", "create_file", {"kind": "csv"}),
        ])),
        _FakeStream(["Fayl tayyor"], _Message("Fayl tayyor")),
    ]
    client = _tool_client(monkeypatch, replies, sent, calls)

    async def on_tool(name, payload):
        calls.append((name, payload))
        return "«a.csv» tayyorlandi."

    reply = asyncio.run(
        client.ask([{"role": "user", "content": "csv qil"}], None, None, on_tool)
    )

    assert reply.text == "Fayl tayyor"
    assert calls == [("create_file", {"kind": "csv"})]
    # Ikkinchi so'rovda qurol natijasi qaytgan bo'lsin.
    result = sent[1][-1]
    assert result["role"] == "user"
    assert result["content"][0]["type"] == "tool_result"
    assert result["content"][0]["tool_use_id"] == "t1"


def test_tool_failure_is_reported_back_to_the_model(monkeypatch):
    import asyncio

    sent, calls = [], []
    replies = [
        _FakeStream([], _Message("", stop_reason="tool_use", content=[
            _ToolUse("t1", "create_file", {"kind": "mp3"}),
        ])),
        _FakeStream(["Uzr, mp3 qila olmayman"], _Message("Uzr, mp3 qila olmayman")),
    ]
    client = _tool_client(monkeypatch, replies, sent, calls)

    async def on_tool(name, payload):
        raise ValueError("mp3 qo'llab-quvvatlanmaydi")

    reply = asyncio.run(
        client.ask([{"role": "user", "content": "mp3 qil"}], None, None, on_tool)
    )

    # Xatolik javobni yo'q qilmaydi — model uni o'qib, tushuntirib beradi.
    assert "mp3" in reply.text
    result = sent[1][-1]["content"][0]
    assert result["is_error"] is True
    assert "mp3" in result["content"]


def test_paused_turn_is_resumed(monkeypatch):
    import asyncio

    pytest.importorskip("anthropic")
    from claude_bot.claude import ClaudeClient

    client = ClaudeClient(
        "sk-test",
        model="claude-opus-5",
        system="test",
        max_tokens=100,
        effort="low",
        use_fallbacks=False,
        tool_mode="web",
    )

    sent = []
    replies = [
        _FakeStream(["qidiryapman"], _Message("qism", stop_reason="pause_turn")),
        _FakeStream([" va javob"], _Message("to'liq javob")),
    ]

    def fake_stream(**kwargs):
        sent.append(kwargs["messages"])
        return replies[len(sent) - 1]

    monkeypatch.setattr(client._client.messages, "stream", fake_stream, raising=False)

    reply = asyncio.run(client.ask([{"role": "user", "content": "hi"}]))

    assert reply.text == "to'liq javob"
    assert len(sent) == 2
    # Ikkinchi so'rovda birinchi javob assistant xabari sifatida qaytadi.
    assert sent[1][-1]["role"] == "assistant"


# --- webhook rejimi ---------------------------------------------------------

def test_webhook_path_hides_the_token():
    pytest.importorskip("aiogram")
    from claude_bot import webhook

    token = "123456:AAHsupersecret"
    path = webhook.secret_path(token)

    assert path.startswith("/telegram/")
    assert token not in path
    assert "AAHsupersecret" not in path
    # Bir xil token — bir xil manzil (qayta deploy'da o'zgarib ketmasin).
    assert webhook.secret_path(token) == path
    assert webhook.secret_path("boshqa:token") != path


def test_webhook_secret_differs_from_the_path():
    pytest.importorskip("aiogram")
    from claude_bot import webhook

    token = "123456:AAHsupersecret"
    assert token not in webhook.secret_token(token)
    assert webhook.secret_token(token) not in webhook.secret_path(token)


def test_webhook_app_serves_health_and_telegram_routes():
    pytest.importorskip("aiogram")
    from aiogram import Bot, Dispatcher

    from claude_bot import webhook

    bot = Bot("123456:AAHfake-token-for-tests")
    app = webhook.build_app(
        bot, Dispatcher(), path="/telegram/abc", secret="s3cret"
    )
    routes = {
        (route.method, route.resource.canonical) for route in app.router.routes()
    }

    assert ("GET", "/") in routes
    assert ("POST", "/telegram/abc") in routes


# --- rasm/fayl handler'i ----------------------------------------------------

def _run_media(reply, message, payload=b"", history=None):
    import asyncio

    pytest.importorskip("aiogram")
    from claude_bot import main

    from claude_bot.session import UsageTracker

    claude = _StubClaude(reply)
    history = history or ChatHistory()
    asyncio.run(
        main.on_media(
            message,
            bot=_StubBot(payload),
            claude=claude,
            history=history,
            limiter=RateLimiter(per_minute=0),
            usage=UsageTracker(),
        )
    )
    return claude, history


def test_photo_reaches_claude_as_an_image_with_the_caption():
    photo = [type("Size", (), {"file_id": "abc", "file_size": 1024})()]
    message = _StubMessage(photo=photo, caption="Bu jadvalda xato bormi?")

    claude, history = _run_media(_reply("Xato yo'q"), message, payload=b"jpeg-bytes")

    content = claude.seen[-1]["content"]
    assert content[0]["type"] == "image"
    assert content[1]["text"] == "Bu jadvalda xato bormi?"
    # Rasm tarixda qoladi — «buni PDF qil» deb davom ettirish uchun kerak.
    # Qachon chiqib ketishi alohida testlarda tekshiriladi.
    assert history.get(1)[0]["content"] == content


def test_photo_without_a_caption_gets_a_default_question():
    from claude_bot import texts

    photo = [type("Size", (), {"file_id": "abc", "file_size": 10})()]
    message = _StubMessage(photo=photo)

    claude, _ = _run_media(_reply("javob"), message, payload=b"x")

    assert claude.seen[-1]["content"][1]["text"] == texts.DEFAULT_IMAGE_PROMPT


def test_pdf_document_becomes_a_document_block():
    document = type("Doc", (), {
        "file_id": "f1", "file_size": 2048,
        "file_name": "hisobot.pdf", "mime_type": "application/pdf",
    })()
    message = _StubMessage(document=document, caption="Xulosa qil")

    claude, _ = _run_media(_reply("xulosa"), message, payload=b"%PDF-1.4")

    assert claude.seen[-1]["content"][0]["type"] == "document"


def test_unsupported_file_is_refused_before_any_download():
    from claude_bot import texts

    document = type("Doc", (), {
        "file_id": "f2", "file_size": 100,
        "file_name": "klip.mp4", "mime_type": "video/mp4",
    })()
    message = _StubMessage(document=document)

    claude, _ = _run_media(_reply("javob"), message)

    assert claude.seen is None
    assert texts.UNSUPPORTED_FILE in message.edits[-1]


def test_oversized_file_is_refused():
    from claude_bot import config

    document = type("Doc", (), {
        "file_id": "f3", "file_size": (config.MAX_FILE_MB + 1) * 1024 * 1024,
        "file_name": "katta.pdf", "mime_type": "application/pdf",
    })()
    message = _StubMessage(document=document)

    claude, _ = _run_media(_reply("javob"), message)

    assert claude.seen is None
    assert "katta" in message.edits[-1]


def test_empty_text_file_is_explained():
    from claude_bot import texts

    document = type("Doc", (), {
        "file_id": "f4", "file_size": 3,
        "file_name": "bosh.txt", "mime_type": "text/plain",
    })()
    message = _StubMessage(document=document)

    claude, _ = _run_media(_reply("javob"), message, payload=b"   ")

    assert claude.seen is None
    assert message.edits[-1] == texts.EMPTY_FILE


def test_pdf_command_sends_a_document_built_from_the_last_answer():
    import asyncio

    pytest.importorskip("aiogram")
    pytest.importorskip("reportlab")
    from claude_bot import main

    history = ChatHistory()
    history.add(1, "user", "ro'yxatni yoz")
    history.add(1, "assistant", "## Qarindoshlar\n\n- Akram — 300 000")

    message = _StubMessage()
    asyncio.run(main.cmd_pdf(message, history=history))

    document = message.outbox[-1]
    assert document.filename.endswith(".pdf")
    assert document.data.startswith(b"%PDF")


def test_pdf_command_without_an_answer_explains_itself():
    import asyncio

    pytest.importorskip("aiogram")
    from claude_bot import main, texts

    message = _StubMessage()
    asyncio.run(main.cmd_pdf(message, history=ChatHistory()))

    assert message.outbox == [texts.NOTHING_TO_EXPORT]


def test_attachment_survives_the_answer_that_follows_it():
    history = ChatHistory()
    blocks = [{"type": "image", "source": {"data": "aaa"}}]

    history.add(1, "user", blocks, marker="[rasm yuborildi]")
    history.add(1, "assistant", "rasmda ro'yxat bor")

    # Rasm haqida yana savol berish mumkin bo'lishi kerak — javob qo'shilgani
    # uni kontekstdan chiqarib yubormasin.
    assert history.get(1)[0]["content"] == blocks


def test_friendly_error_maps_known_failures():
    anthropic = pytest.importorskip("anthropic")
    from claude_bot import texts
    from claude_bot.claude import friendly_error

    assert friendly_error(
        anthropic.APIConnectionError(request=None)
    ) == texts.ERR_NETWORK
    assert friendly_error(RuntimeError("?")) == texts.ERR_UNKNOWN
